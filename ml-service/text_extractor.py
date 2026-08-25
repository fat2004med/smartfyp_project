import io
import os
from PyPDF2 import PdfReader
from docx import Document

def extract_text_from_file(file_storage, filename: str) -> str:
    """
    Extracts raw UTF-8 clean text from uploaded PDF, DOCX, or TXT file storage objects.
    """
    if not filename:
        return ""

    ext = os.path.splitext(filename)[1].lower()
    text = ""

    try:
        if ext == ".pdf":
            # PDF Extraction
            reader = PdfReader(file_storage)
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_content = page.extract_text()
                if page_content:
                    pages_text.append(page_content)
            text = "\n".join(pages_text)

        elif ext in [".docx", ".doc"]:
            # Word Document Extraction
            doc = Document(file_storage)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs)

        elif ext in [".txt", ".md", ".csv", ".json"]:
            # Plain Text Extraction
            raw_bytes = file_storage.read()
            try:
                text = raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = raw_bytes.decode("latin-1", errors="ignore")
        else:
            # General fallback
            raw_bytes = file_storage.read()
            text = raw_bytes.decode("utf-8", errors="ignore")

    except Exception as e:
        print(f"❌ [TextExtractor] Error reading {filename}: {e}")
        raise ValueError(f"Failed to parse file '{filename}': {str(e)}")

    # Clean whitespace and standardize line breaks
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return cleaned
